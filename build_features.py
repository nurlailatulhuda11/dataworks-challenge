import glob
import pandas as pd
from openpyxl import load_workbook

KEEP = [
    "Pore Volume (PORV)", "Porosity (PORO)", "Permeability I (PERMX)",
    "Permeability K (PERMZ)", "Transmissibility I (TRANX)",
    "Connate Water Saturation (SWL)", "Critical Water Saturation (SWCR)",
    "Critical Gas Saturation (SGCR)", "Residual Oil to Water (SOWCR)",
    "Pressure (PRESSURE)", "Water Saturation (SWAT)", "Oil Saturation (SOIL)",
    "Fluid In Place Region (FIPNUM)",
]

rows = []
for path in sorted(glob.glob("data/reservoir_cases/0[3-8]*.xlsx")):
    print("Membuka", path, flush=True)
    wb = load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = next(it)
        df = pd.DataFrame(it, columns=header)
        df = df[df["Pore Volume (PORV)"] > 0]

        f = {"case": int(ws.title.split()[-1])}
        f["n_active"] = len(df)
        f["porv_sum"] = df["Pore Volume (PORV)"].sum()
        f["oil_in_place"] = (df["Pore Volume (PORV)"] * df["Oil Saturation (SOIL)"]).sum()
        for c in KEEP[1:]:
            f[c + "_mean"] = df[c].mean()
        f["PERMX_median"] = df["Permeability I (PERMX)"].median()
        f["PERMX_p90"] = df["Permeability I (PERMX)"].quantile(0.9)
        f["frac_FIPNUM2"] = (df["Fluid In Place Region (FIPNUM)"] == 2).mean()

        rows.append(f)
        print("  selesai Case", f["case"], flush=True)
    wb.close()

    # simpan hasil separa selepas setiap fail
    pd.DataFrame(rows).sort_values("case").to_csv("data/case_features.csv", index=False)

print("Siap:", len(rows), "kes")