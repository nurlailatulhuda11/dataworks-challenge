import glob, os
import pandas as pd
from openpyxl import load_workbook

os.makedirs("data/reservoir_parquet", exist_ok=True)

for path in sorted(glob.glob("data/reservoir_cases/*.xlsx")):
    base = os.path.splitext(os.path.basename(path))[0].strip().replace(" ", "_").lower()
    wb = load_workbook(path, read_only=True)

    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        df = pd.DataFrame(rows, columns=header)
        out = f"data/reservoir_parquet/{base}__{ws.title.strip().replace(' ', '_').lower()}.parquet"
        df.to_parquet(out, index=False)
        print("Disimpan:", out, df.shape)

    wb.close()