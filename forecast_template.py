from datetime import datetime
import numpy as np
import pandas as pd
from openpyxl import load_workbook

PHASES = {"Gas": "gas_production", "Oil": "oil_production", "Water": "water_production"}
N_FIT = 8
LO, HI = 0.90, 1.03          # had perubahan kadar setiap suku tahun


def load(name):
    return pd.read_csv(f"data/quarterly_{name}.csv", parse_dates=["Date"], index_col="Date")


def forecast_case(rate_hist, cum_last, last_date, new_dates):
    r = np.asarray(rate_hist, dtype=float)[-N_FIT:]
    if len(r) >= 3 and r.min() > 0:
        slope = np.polyfit(np.arange(len(r)), np.log(r), 1)[0]
        g = float(np.clip(np.exp(slope), LO, HI))
    else:
        g = 1.0
    k = np.arange(1, len(new_dates) + 1)
    rates = r[-1] * g ** k
    all_rates = np.concatenate([[r[-1]], rates])
    all_dates = pd.DatetimeIndex([last_date]).append(pd.DatetimeIndex(new_dates))
    days = np.diff(all_dates.values).astype("timedelta64[D]").astype(float)
    cum = cum_last + np.cumsum(0.5 * (all_rates[:-1] + all_rates[1:]) * days)
    return cum


rate = {p: load(f"{f}_rate") for p, f in PHASES.items()}
cumu = {p: load(f"{f}_cumulative") for p, f in PHASES.items()}
dates = rate["Oil"].index

# ---------- Backtest: potong data pada 8 suku tahun sebelum akhir ----------
cut = len(dates) - 1 - 8
print("Backtest (data dipotong pada", dates[cut].date(), ", ramal 8 suku tahun):")
for p in PHASES:
    errs = []
    for c in rate[p].columns:
        pred = forecast_case(rate[p][c].values[: cut + 1], cumu[p][c].iloc[cut],
                             dates[cut], dates[cut + 1:])[-1]
        true = cumu[p][c].iloc[-1]
        if true > 0:
            errs.append(abs(pred - true) / true)
    print(f"  {p}: ralat peratus median = {np.median(errs):.2%}, min-maks = "
          f"{np.min(errs):.2%} - {np.max(errs):.2%}")

# ---------- Isi template ----------
wb = load_workbook("09 Template Deliverable.xlsx")
last_date = dates[-1]
for ws in wb.worksheets:
    case = int(ws.title.split()[-1])
    col_name = f"Case_{case}"
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or "")
        for p in PHASES:
            if h.startswith(p):
                cols[p] = c
    new_dates = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None:
            break
        new_dates.append(v if isinstance(v, datetime) else datetime.strptime(str(v), "%m/%d/%Y"))
    for p, c in cols.items():
        cum = forecast_case(rate[p][col_name].values, cumu[p][col_name].iloc[-1],
                            last_date, new_dates)
        for i, val in enumerate(cum):
            ws.cell(2 + i, c).value = float(val)

import os
os.makedirs("outputs", exist_ok=True)
wb.save("outputs/Deliverable_filled.xlsx")
print("Siap: outputs/Deliverable_filled.xlsx,", len(wb.worksheets), "worksheet")
