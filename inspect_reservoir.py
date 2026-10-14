import glob
from openpyxl import load_workbook

for path in sorted(glob.glob("data/reservoir_cases/*.xlsx")):
    wb = load_workbook(path, read_only=True)
    print(f"\n=== {path} ===")
    for ws in wb.worksheets:
        print(f"  Worksheet: {ws.title} | dimensi: {ws.max_row} baris x {ws.max_column} lajur")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            print("   ", row[:8])
            if i == 2:
                break
    wb.close()